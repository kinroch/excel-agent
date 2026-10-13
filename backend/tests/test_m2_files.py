# -*- coding: utf-8 -*-
"""
M2 数据接入与文件解析 · 验收测试

覆盖范围（对齐 PRD 4.2 验收标准与技术文档 5.5 F1~F8）
------------------------------------------------------
    AC-4.2-1  上传合法 xlsx → 完成解析 → 可展示数据预览
    AC-4.2-2  多 Sheet 文件 → 列出全部 Sheet 并支持切换
    AC-4.2-3  超大 / 格式不支持 / 内容损坏 → 拒绝或失败，并给出**具体原因**
    AC-4.2-4  含合并单元格 → 正确展开，不出现错位或空列
    AC-4.2-5  含缺失值 → 字段说明标注缺失率 + 提供处理选项（F8）
    AC-4.2-6  含编码差异的 csv → 自动识别编码并正确展示中文
    F-4.2-1/2  上传受理（202 语义）与格式/大小/空文件校验
    F-4.2-3/4  多 Sheet 识别、表头识别（标题行/多级表头）、数据类型推断
    F-4.2-7    预览分页（含越界页码回落）
    F-4.2-9    文件列表与删除
    F-4.7-2    敏感字段识别与示例值脱敏
    技术文档 2.2 / AC-4.1-4  数据隔离：看不到、也删不掉他人的文件
    技术文档 3.6  崩溃恢复（残留 parsing → failed）

运行方式
--------
    cd backend
    ../../.venv/Scripts/python.exe -m pytest tests/test_m2_files.py -q
"""

import io
import uuid

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.core.config import settings
from app.infra.db import SessionLocal
from app.infra.models import File, utc_now_iso
from app.main import app
from app.services import file_service, parse_service

# ============================================================
# 夹具与工具
# ============================================================


@pytest.fixture(scope="module")
def client():
    """模块级 TestClient（with 语法触发 lifespan：建库 + 崩溃恢复）。

    产出：
        TestClient: 测试客户端

    说明：
        Starlette 的 TestClient 会在应用调用返回前执行完 BackgroundTasks，
        因此 `client.post("/api/v1/files", ...)` 返回时解析已经完成——
        测试无需真的轮询等待。若需验证「轮询」语义，仍可直接调 F3 断言终态。
    """
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture(scope="module")
def auth(client: TestClient) -> dict[str, str]:
    """注册并登录一个普通用户，返回带令牌的请求头。

    产出：
        dict[str, str]: {"Authorization": "Bearer <token>"}
    """
    username = f"m2_{uuid.uuid4().hex[:12]}@example.com"
    password = "Passw0rd123"
    assert client.post(
        "/api/v1/auth/register", json={"username": username, "password": password}
    ).status_code == 201
    login = client.post(
        "/api/v1/auth/login", json={"username": username, "password": password, "remember": False}
    )
    assert login.status_code == 200, login.text
    return {"Authorization": f"Bearer {login.json()['data']['access_token']}"}


@pytest.fixture(scope="module")
def other_auth(client: TestClient) -> dict[str, str]:
    """第二个用户（用于验证跨用户数据隔离）。

    产出：
        dict[str, str]: {"Authorization": "Bearer <token>"}
    """
    username = f"m2o_{uuid.uuid4().hex[:12]}@example.com"
    password = "Passw0rd123"
    client.post("/api/v1/auth/register", json={"username": username, "password": password})
    login = client.post(
        "/api/v1/auth/login", json={"username": username, "password": password, "remember": False}
    )
    return {"Authorization": f"Bearer {login.json()['data']['access_token']}"}


def _xlsx_bytes() -> bytes:
    """构造样本 xlsx（覆盖 M2 全部解析要点）。

    结构（行号按 Excel 从 1 起）：
        第 1 行  合并大标题「2026 年第一季度销售明细表」（A1:F1）→ 应被判为说明性内容
        第 2 行  多级表头上级：A2:B2「销售维度」、C2:F2「指标」
        第 3 行  表头：部门/负责人/金额/数量/手机号/邮箱
        第 4-6 行 数据；其中 A5:A6 纵向合并（测向下填充）、D5 缺失
    另有第二个 Sheet「部门对照」，第三个 Sheet「空表」（应被跳过）。

    返回：
        bytes: xlsx 文件字节
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "销售明细"

    sheet["A1"] = "2026 年第一季度销售明细表"
    sheet.merge_cells("A1:F1")
    sheet["A2"] = "销售维度"
    sheet.merge_cells("A2:B2")
    sheet["C2"] = "指标"
    sheet.merge_cells("C2:F2")
    for column, name in zip("ABCDEF", ["部门", "负责人", "金额", "数量", "手机号", "邮箱"], strict=True):
        sheet[f"{column}3"] = name

    sheet["A4"], sheet["B4"] = "华东", "张三"
    sheet["C4"], sheet["D4"] = 1200.5, 3
    sheet["E4"], sheet["F4"] = "13812345678", "zhangsan@example.com"

    sheet["A5"], sheet["B5"] = "华南", "李四"
    sheet["C5"] = 900  # D5 故意留空（缺失值）
    sheet["E5"], sheet["F5"] = "13900001111", "lisi@example.com"
    sheet.merge_cells("A5:A6")  # 数据区纵向合并：A6 应被填充为「华南」

    sheet["B6"], sheet["C6"], sheet["D6"] = "王五", 800.25, 5
    sheet["E6"], sheet["F6"] = "13711112222", "wangwu@example.com"

    second = workbook.create_sheet("部门对照")
    second["A1"], second["B1"] = "部门", "编码"
    second["A2"], second["B2"] = "华东", "E01"
    second["A3"], second["B3"] = "华南", "S01"

    workbook.create_sheet("空表")  # 空 Sheet：应被跳过，不产生结构

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _gbk_csv_bytes() -> bytes:
    """构造 GBK 编码的 csv（AC-4.2-6：中文需正确识别）。

    返回：
        bytes: GBK 编码的 csv 字节
    """
    text = "部门,负责人,金额,电话\n华东,张三,1200.5,13812345678\n华南,李四,900,13900001111\n华北,王五,,13711112222\n"
    return text.encode("gbk")


def _upload(client: TestClient, headers: dict[str, str], name: str, payload: bytes, mime: str):
    """调用 F1 上传接口。

    返回：
        httpx.Response: 原始响应
    """
    return client.post("/api/v1/files", headers=headers, files={"file": (name, payload, mime)})


def _upload_ok(client: TestClient, headers: dict[str, str], name: str, payload: bytes, mime: str) -> int:
    """上传并断言受理成功，返回 file_id。"""
    response = _upload(client, headers, name, payload, mime)
    assert response.status_code == 202, response.text
    body = response.json()
    assert body["code"] == 0
    assert body["data"]["parse_status"] == "parsing"  # 受理即返回 parsing（技术文档 5.5 F1）
    return body["data"]["file_id"]


@pytest.fixture(scope="module")
def xlsx_file_id(client: TestClient, auth: dict[str, str]) -> int:
    """上传一次样本 xlsx 并返回其 file_id（供多个用例复用，避免重复解析）。

    产出：
        int: file_id（解析已完成）
    """
    file_id = _upload_ok(client, auth, "sales.xlsx", _xlsx_bytes(), XLSX_MIME)
    detail = client.get(f"/api/v1/files/{file_id}", headers=auth).json()["data"]
    assert detail["parse_status"] == "parsed", detail
    return file_id


# ============================================================
# 一、F1 上传受理与校验
# ============================================================


def test_f1_accept_and_parse_xlsx(client: TestClient, auth: dict[str, str]):
    """F1 + AC-4.2-1：上传 xlsx 受理返回 202，随后可拿到解析完成的项目。"""
    response = _upload(client, auth, "sales.xlsx", _xlsx_bytes(), XLSX_MIME)
    assert response.status_code == 202, response.text
    body = response.json()
    assert body["code"] == 0
    assert isinstance(body["data"]["file_id"], int)
    assert body["data"]["original_name"] == "sales.xlsx"

    detail = client.get(f"/api/v1/files/{body['data']['file_id']}", headers=auth).json()["data"]
    assert detail["parse_status"] == "parsed"
    assert detail["parse_progress"] == 100
    assert detail["error"] is None
    assert detail["stage_text"] == "解析完成"
    # 质量小结五项（UI/UX 6.3 stat-strip）
    assert detail["quality"]["row_total"] == 5  # 3 行 + 2 行（两个 Sheet）
    assert detail["quality"]["merged_cells"] >= 4


def test_f1_reject_unsupported_format(client: TestClient, auth: dict[str, str]):
    """F1/AC-4.2-3：格式不支持 → 2001，且文案点明支持范围。"""
    response = _upload(client, auth, "report.pdf", b"%PDF-1.4 fake", "application/pdf")
    assert response.status_code == 400
    body = response.json()
    assert body["code"] == 2001
    assert "xlsx" in body["message"] and "csv" in body["message"]


def test_f1_reject_mime_contradiction(client: TestClient, auth: dict[str, str]):
    """F1/AC-4.2-3：扩展名与 MIME 明确矛盾 → 拒绝（防伪装扩展名）。"""
    response = _upload(client, auth, "fake.csv", b"a,b\n1,2\n", "image/png")
    assert response.status_code == 400
    assert response.json()["code"] == 2001


def test_f1_reject_empty_file(client: TestClient, auth: dict[str, str]):
    """F1：空文件 → 2003。"""
    response = _upload(client, auth, "empty.csv", b"", "text/csv")
    assert response.status_code == 400
    assert response.json()["code"] == 2003


def test_f1_reject_oversize(client: TestClient, auth: dict[str, str]):
    """F1/AC-4.2-3：超出大小上限 → 2002（临时把上限降到 1MB，避免生成 50MB 测试数据）。"""
    original = settings.max_upload_size_mb
    settings.max_upload_size_mb = 1
    try:
        payload = b"a,b\n" + b"1,2\n" * 400_000  # 约 1.6MB
        response = _upload(client, auth, "big.csv", payload, "text/csv")
        assert response.status_code == 400, response.text
        body = response.json()
        assert body["code"] == 2002
        assert "1MB" in body["message"]  # 文案与上限设置一致（不写死 50MB）
    finally:
        settings.max_upload_size_mb = original


def test_f1_reject_encrypted_xlsx(client: TestClient, auth: dict[str, str]):
    """F1：加密的 xlsx（OLE/CFB 容器）→ 2004，原因明确为「加密受保护」。"""
    payload = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 200
    response = _upload(client, auth, "locked.xlsx", payload, XLSX_MIME)
    assert response.status_code == 422, response.text
    body = response.json()
    assert body["code"] == 2004
    assert "加密" in body["message"]


def test_f1_corrupt_xlsx_marks_failed(client: TestClient, auth: dict[str, str]):
    """F1+F3/AC-4.2-3：内容损坏的 xlsx 受理后解析失败，F3 给出可读原因。"""
    file_id = _upload_ok(client, auth, "broken.xlsx", b"this is not a zip file", XLSX_MIME)
    detail = client.get(f"/api/v1/files/{file_id}", headers=auth).json()["data"]
    assert detail["parse_status"] == "failed"
    assert detail["error"] and "损坏" in detail["error"]
    assert detail["parse_progress"] == 0
    # 解析失败后不应留下可读的 Sheet 结构
    sheets = client.get(f"/api/v1/files/{file_id}/sheets", headers=auth)
    assert sheets.status_code == 422
    assert sheets.json()["code"] == 2004


# ============================================================
# 二、F5/F6/F7 解析结果（AC-4.2-2/4/5）
# ============================================================


def test_f5_multi_sheet_and_notes(client: TestClient, auth: dict[str, str], xlsx_file_id: int):
    """AC-4.2-2 / AC-4.2-4 / AC-4.2-5：多 Sheet 列出 + 解析处理信息完整。"""
    body = client.get(f"/api/v1/files/{xlsx_file_id}/sheets", headers=auth).json()
    assert body["code"] == 0
    data = body["data"]
    assert data["sheet_count"] == 2  # 「空表」被跳过，只保留 2 个有效 Sheet

    first, second = data["sheets"]
    assert first["index"] == 0 and first["name"] == "销售明细"
    assert second["index"] == 1 and second["name"] == "部门对照"
    assert first["row_count"] == 3 and first["col_count"] == 6
    assert first["header_row"] == 2  # 第 1 行是合并大标题，表头在第 2 行（多级表头起始行）

    note = first["header_note"]
    assert "表头识别在第 2 行" in note
    assert "多级表头" in note                  # 上级「销售维度/指标」已降级合并
    assert "合并单元格已向下填充展开" in note   # AC-4.2-4


def test_f5_drop_blank_rows_and_columns(client: TestClient, auth: dict[str, str]):
    """AC-4.2-4 边界：全空行/全空列被剔除并如实告知（避免出现「列3」这类空字段）。"""
    # 第 3 列全空、第 3 个数据行是全空行
    payload = "部门,金额,,备注\n华东,100,,a\n\n华南,200,,b\n".encode()
    file_id = _upload_ok(client, auth, "blank.csv", payload, "text/csv")

    sheets = client.get(f"/api/v1/files/{file_id}/sheets", headers=auth).json()["data"]["sheets"]
    sheet = sheets[0]
    assert sheet["col_count"] == 3       # 全空列已剔除（4 列 → 3 列）
    assert sheet["row_count"] == 2       # 全空行已剔除（3 行 → 2 行）
    assert "已忽略 1 个全空行、1 个全空列" in sheet["header_note"]

    preview = client.get(f"/api/v1/files/{file_id}/preview?sheet=0", headers=auth).json()["data"]
    assert preview["columns"] == ["部门", "金额", "备注"]
    assert preview["rows"] == [["华东", "100", "a"], ["华南", "200", "b"]]


def test_f6_fields_types_and_missing(client: TestClient, auth: dict[str, str], xlsx_file_id: int):
    """F-4.2-4/8 + AC-4.2-5：字段类型推断、缺失率、推断依据、高缺失阈值。"""
    body = client.get(f"/api/v1/files/{xlsx_file_id}/fields?sheet=0", headers=auth).json()
    assert body["code"] == 0
    data = body["data"]
    assert data["sheet_index"] == 0 and data["sheet_name"] == "销售明细"
    assert data["high_missing_threshold"] == pytest.approx(parse_service.HIGH_MISSING_THRESHOLD)

    by_name = {item["name"]: item for item in data["fields"]}
    # 多级表头降级为「上级/下级」
    assert "销售维度/部门" in by_name and "指标/金额" in by_name

    assert by_name["指标/金额"]["dtype"] == "小数"
    assert by_name["指标/数量"]["dtype"] == "整数"
    assert by_name["销售维度/部门"]["dtype"] == "字符串"

    # 缺失率：3 行数据中「数量」缺 1 个 → 1/3
    assert by_name["指标/数量"]["missing_rate"] == pytest.approx(0.3333, abs=1e-4)
    assert by_name["指标/数量"]["samples"] == ["3", "5"]

    # 每个字段都必须给出推断依据（F6 契约要求 infer_basis 非空）
    assert all(item["infer_basis"] for item in data["fields"])
    assert len(data["fields"]) == 6


def test_f6_sensitive_masking(client: TestClient, auth: dict[str, str], xlsx_file_id: int):
    """F-4.7-2/AC-4.7-2：敏感字段被标记，示例值已脱敏（不下发原始值）。"""
    data = client.get(f"/api/v1/files/{xlsx_file_id}/fields?sheet=0", headers=auth).json()["data"]
    by_name = {item["name"]: item for item in data["fields"]}

    phone = by_name["指标/手机号"]
    assert phone["sensitive"] is True
    assert phone["sensitive_kind"] == "手机号"
    assert phone["dtype"] == "字符串"  # 标识类字段不得参与数值统计
    assert "138****5678" in phone["samples"]
    assert "13812345678" not in phone["samples"]  # 原始值绝不出现

    email = by_name["指标/邮箱"]
    assert email["sensitive"] is True
    assert email["sensitive_kind"] == "邮箱"
    assert all("***" in s for s in email["samples"])

    # 非敏感字段不受影响
    assert by_name["销售维度/部门"]["sensitive"] is False
    assert by_name["销售维度/部门"]["sensitive_kind"] is None


def test_f7_preview_pagination(client: TestClient, auth: dict[str, str], xlsx_file_id: int):
    """F-4.2-7/AC-4.2-1：预览分页、缺失值以 null 返回、越界页码回落末页。"""
    page1 = client.get(
        f"/api/v1/files/{xlsx_file_id}/preview?sheet=0&page=1&size=2", headers=auth
    ).json()["data"]
    assert page1["page"] == 1 and page1["per_page"] == 2
    assert page1["total_rows"] == 3 and page1["total_pages"] == 2
    assert len(page1["rows"]) == 2
    assert page1["columns"][0] == "销售维度/部门"
    assert page1["rows"][0][0] == "华东"
    # AC-4.2-5：缺失单元格以 null 下发，前端渲染「（缺失）」
    assert page1["rows"][1][3] is None

    page2 = client.get(
        f"/api/v1/files/{xlsx_file_id}/preview?sheet=0&page=2&size=2", headers=auth
    ).json()["data"]
    assert page2["page"] == 2 and len(page2["rows"]) == 1
    # AC-4.2-4：合并单元格已向下填充（A5:A6 → 第 3 行部门为「华南」，不是 None）
    assert page2["rows"][0][0] == "华南"

    # 越界页码回落到末页（不报错）
    overflow = client.get(
        f"/api/v1/files/{xlsx_file_id}/preview?sheet=0&page=99&size=2", headers=auth
    ).json()["data"]
    assert overflow["page"] == 2
    assert len(overflow["rows"]) == 1


def test_f7_preview_second_sheet(client: TestClient, auth: dict[str, str], xlsx_file_id: int):
    """AC-4.2-2：Sheet 切换后预览的是对应 Sheet 的数据。"""
    data = client.get(f"/api/v1/files/{xlsx_file_id}/preview?sheet=1", headers=auth).json()["data"]
    assert data["sheet_index"] == 1
    assert data["sheet_name"] == "部门对照"
    assert data["columns"] == ["部门", "编码"]
    assert data["rows"] == [["华东", "E01"], ["华南", "S01"]]


def test_f7_preview_invalid_sheet(client: TestClient, auth: dict[str, str], xlsx_file_id: int):
    """F7：Sheet 索引越界 → 2004（解析失败段，语义为「该工作表不存在」）。"""
    response = client.get(f"/api/v1/files/{xlsx_file_id}/preview?sheet=99", headers=auth)
    assert response.status_code == 422
    assert response.json()["code"] == 2004


def test_ac_4_2_6_gbk_csv_encoding(client: TestClient, auth: dict[str, str]):
    """AC-4.2-6：GBK 编码的 csv 自动识别，中文正确展示。"""
    file_id = _upload_ok(client, auth, "dept.csv", _gbk_csv_bytes(), "text/csv")

    detail = client.get(f"/api/v1/files/{file_id}", headers=auth).json()["data"]
    assert detail["parse_status"] == "parsed"
    # GBK 是 GB18030 的子集，识别结果可为 gb18030 / gbk 任一（均能正确解码）
    assert detail["encoding"] in {"gb18030", "gbk"}

    preview = client.get(f"/api/v1/files/{file_id}/preview?sheet=0", headers=auth).json()["data"]
    assert preview["rows"][0][0] == "华东"
    assert preview["rows"][0][1] == "张三"

    sheets = client.get(f"/api/v1/files/{file_id}/sheets", headers=auth).json()["data"]["sheets"]
    assert "编码自动识别为" in sheets[0]["header_note"]
    assert sheets[0]["name"] == "dept"  # csv 的 Sheet 名取文件名主干

    fields = client.get(f"/api/v1/files/{file_id}/fields?sheet=0", headers=auth).json()["data"]["fields"]
    by_name = {item["name"]: item for item in fields}
    assert by_name["部门"]["samples"][0] == "华东"
    assert by_name["电话"]["sensitive"] is True


def test_csv_utf8_bom(client: TestClient, auth: dict[str, str]):
    """AC-4.2-6 补充：带 UTF-8 BOM 的 csv（Excel 另存为的常见形态）不出现 \\ufeff 脏字符。"""
    payload = "\ufeff部门,金额\n华东,100\n".encode()
    file_id = _upload_ok(client, auth, "bom.csv", payload, "text/csv")
    preview = client.get(f"/api/v1/files/{file_id}/preview?sheet=0", headers=auth).json()["data"]
    assert preview["columns"] == ["部门", "金额"]
    assert preview["rows"][0][0] == "华东"


def test_csv_utf8_chinese_not_misdetected(client: TestClient, auth: dict[str, str]):
    """AC-4.2-6 回归：**无 BOM 的 UTF-8 中文 csv 不得被误判为 GB18030**。

    为什么单列此用例
    ----------------
    实测踩到的严重缺陷：若按「解出中文字符数最多」选编码，GB18030 每 2 字节产 1 个
    汉字、UTF-8 每 3 字节产 1 个，同一段字节用 GB18030 解出的汉字数总是更多，
    导致**所有 UTF-8 中文 csv 都被误判**，中文变乱码（`部门` → `閮ㄩ棬`）。
    本用例锁定该回归：断言编码识别为 utf-8 且中文原样正确。
    """
    payload = "部门,负责人,金额\n华东,张三,1200.5\n华南,李四,900\n".encode()
    file_id = _upload_ok(client, auth, "utf8_cn.csv", payload, "text/csv")

    detail = client.get(f"/api/v1/files/{file_id}", headers=auth).json()["data"]
    assert detail["encoding"] == "utf-8"

    preview = client.get(f"/api/v1/files/{file_id}/preview?sheet=0", headers=auth).json()["data"]
    assert preview["columns"] == ["部门", "负责人", "金额"]
    assert preview["rows"][0] == ["华东", "张三", "1200.5"]


# ============================================================
# 三、F8 缺失值策略（AC-4.2-5）
# ============================================================


def test_f8_missing_policy_switch(client: TestClient, auth: dict[str, str], xlsx_file_id: int):
    """AC-4.2-5：缺失值处理策略可切换（提示/填充/剔除），且不破坏原始数据。"""
    for policy in ("fill", "drop", "prompt"):
        response = client.patch(
            f"/api/v1/files/{xlsx_file_id}/missing-policy", headers=auth, json={"policy": policy}
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["code"] == 0
        assert body["data"] == {"file_id": xlsx_file_id, "missing_policy": policy}

        detail = client.get(f"/api/v1/files/{xlsx_file_id}", headers=auth).json()["data"]
        assert detail["missing_policy"] == policy

    # 策略切换不改变数据本身（预览内容保持一致）——避免「切换策略即改写原件」的隐患
    preview = client.get(f"/api/v1/files/{xlsx_file_id}/preview?sheet=0", headers=auth).json()["data"]
    assert preview["total_rows"] == 3
    assert preview["rows"][1][3] is None  # 缺失值仍在（未被就地填充）


def test_f8_invalid_policy_rejected(client: TestClient, auth: dict[str, str], xlsx_file_id: int):
    """F8：非法策略值由协议层拦截（Pydantic 枚举），不落库、不入业务层。"""
    response = client.patch(
        f"/api/v1/files/{xlsx_file_id}/missing-policy", headers=auth, json={"policy": "unknown"}
    )
    assert response.status_code == 400
    # Pydantic 校验失败由全局处理器转为 9000 + HTTP 400（main.py handle_validation_error）
    assert response.json()["code"] == 9000


# ============================================================
# 四、F2 列表 / F4 删除 / 数据隔离
# ============================================================


def test_f2_file_list_counts(client: TestClient, auth: dict[str, str]):
    """F-4.2-9：列表返回 items/total/analyzable_count，且按时间倒序。"""
    data = client.get("/api/v1/files", headers=auth).json()["data"]
    assert data["total"] == len(data["items"])  # 用例量远小于 FILE_LIST_LIMIT
    assert data["analyzable_count"] >= 1

    created = [item["created_at"] for item in data["items"]]
    assert created == sorted(created, reverse=True)  # 倒序

    first = data["items"][0]
    for key in (
        "id", "original_name", "format", "size_bytes",
        "parse_status", "parse_progress", "sheet_count", "created_at",
    ):
        assert key in first
    # sheet_count 来自 data_sheets 聚合（不是 0 常量）
    assert max(item["sheet_count"] for item in data["items"]) >= 2


def test_f4_delete_then_inaccessible(client: TestClient, auth: dict[str, str]):
    """F-4.2-9：删除后文件在列表与详情中都不再可见。"""
    file_id = _upload_ok(client, auth, "to_delete.csv", "部门,金额\n华东,1\n".encode(), "text/csv")

    deleted = client.delete(f"/api/v1/files/{file_id}", headers=auth)
    assert deleted.status_code == 200, deleted.text
    assert deleted.json()["data"] == {"deleted": True, "file_id": file_id}

    assert client.get(f"/api/v1/files/{file_id}", headers=auth).status_code == 403  # 2005 → 403
    listed = client.get("/api/v1/files", headers=auth).json()["data"]
    assert all(item["id"] != file_id for item in listed["items"])
    # 重复删除同样返回 2005（幂等地不可见）
    assert client.delete(f"/api/v1/files/{file_id}", headers=auth).json()["code"] == 2005


def test_isolation_between_users(
    client: TestClient, auth: dict[str, str], other_auth: dict[str, str], xlsx_file_id: int
):
    """技术文档 2.2：他人文件不可见、不可预览、不可删除（统一 2005，不泄露存在性）。"""
    detail = client.get(f"/api/v1/files/{xlsx_file_id}", headers=other_auth)
    assert detail.status_code == 403 and detail.json()["code"] == 2005
    assert client.get(f"/api/v1/files/{xlsx_file_id}/preview", headers=other_auth).json()["code"] == 2005
    assert client.get(f"/api/v1/files/{xlsx_file_id}/fields", headers=other_auth).json()["code"] == 2005
    assert client.patch(
        f"/api/v1/files/{xlsx_file_id}/missing-policy", headers=other_auth, json={"policy": "drop"}
    ).json()["code"] == 2005
    assert client.delete(f"/api/v1/files/{xlsx_file_id}", headers=other_auth).json()["code"] == 2005

    # 他人的列表里不应出现该文件
    listed = client.get("/api/v1/files", headers=other_auth).json()["data"]
    assert all(item["id"] != xlsx_file_id for item in listed["items"])

    # 原主人仍可正常访问（确认上面的拒绝没有误伤）
    assert client.get(f"/api/v1/files/{xlsx_file_id}", headers=auth).json()["code"] == 0


def test_unauthenticated_requires_login(client: TestClient):
    """技术文档 5.1：全部文件接口需要登录，未带令牌一律 1003（HTTP 401）。"""
    assert client.get("/api/v1/files").status_code == 401
    assert client.get("/api/v1/files").json()["code"] == 1003
    assert client.get("/api/v1/files/1").json()["code"] == 1003
    assert client.delete("/api/v1/files/1").json()["code"] == 1003


# ============================================================
# 五、崩溃恢复（技术文档 3.6）
# ============================================================


def test_fail_stale_parsing_recovers_interrupted(client: TestClient, auth: dict[str, str]):
    """技术文档 3.6：残留的 parsing 记录被置为 failed，并给出可操作提示。"""
    db = SessionLocal()
    try:
        user_id = db.execute(select(File.user_id).limit(1)).scalar_one()
        record = File(
            user_id=user_id,
            original_name="ghost.xlsx",
            storage_path="uploads/0/0_ghost.xlsx",
            format="xlsx",
            size_bytes=10,
            parse_status="parsing",
            parse_progress=40,
            parse_error=None,
            missing_policy="prompt",
            quality_json=None,
            created_at=utc_now_iso(),
            deleted_at=None,
        )
        db.add(record)
        db.commit()
        ghost_id = record.id
    finally:
        db.close()

    # 模拟服务重启
    assert file_service.fail_stale_parsing() >= 1

    db = SessionLocal()
    try:
        ghost = db.get(File, ghost_id)
        assert ghost is not None
        assert ghost.parse_status == "failed"
        assert ghost.parse_error and "重启" in ghost.parse_error
        ghost.deleted_at = utc_now_iso()  # 清理：软删，避免影响列表断言
        db.commit()
    finally:
        db.close()
